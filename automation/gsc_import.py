#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Lit les exports Search Console et en tire un resume que le controle de sante sait
relire sans dependance.

Pourquoi deux etapes. Le controle de sante tourne tous les jours et n'a qu'une
dependance optionnelle, Pillow. Lui faire ouvrir un classeur Excel lui en
ajouterait une obligatoire, openpyxl, pour une donnee qui ne change qu'une fois
par semaine. Ce script-ci porte donc seul la dependance : on le lance a la main
apres avoir telecharge les exports, il ecrit automation/gsc-index.json, et le
controle de sante ne lit plus que ce JSON.

Usage :
    python automation/gsc_import.py                 # lit automation/gsc/
    python automation/gsc_import.py <dossier>       # lit un autre dossier

Les deux exports attendus, tels que la Search Console les nomme :
    <domaine>-Coverage-<date>.xlsx                  (Indexation des pages)
    <domaine>-Performance-on-Search-<date>.xlsx     (Performances, 3 derniers mois)

Le plus recent de chaque type gagne. Les classeurs bruts ne sont pas versionnes
(.gitignore) : seul le resume l'est, il pese quelques kilo-octets et se relit
dans une revue de code.
"""
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl manquant. Installer avec : pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
DEFAUT = ROOT / "automation" / "gsc"
SORTIE = ROOT / "automation" / "gsc-index.json"
BASE = "https://mybusinessnotebook.com"


def sans_accent(s):
    """Les onglets s'appellent « Requetes », « Problemes critiques »... et l'accent
    ne survit pas toujours au passage par le systeme de fichiers. On compare donc
    sur une forme sans accent et en minuscules."""
    return "".join(
        c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn"
    ).lower()


def onglet(wb, prefixe):
    for nom in wb.sheetnames:
        if sans_accent(nom).startswith(sans_accent(prefixe)):
            return wb[nom]
    return None


def lignes(ws):
    """Les lignes de donnees, en-tete retiree, lignes vides ignorees."""
    if ws is None:
        return []
    return [r for r in list(ws.iter_rows(values_only=True))[1:] if r and r[0] is not None]


def plus_recent(dossier, motif):
    trouves = sorted(dossier.glob(motif))
    return trouves[-1] if trouves else None


def date_du_nom(chemin):
    m = re.search(r"(\d{4}-\d{2}-\d{2})", chemin.name)
    return m.group(1) if m else None


def chemin_du_site(url):
    """https://mybusinessnotebook.com/en/x.html -> en/x.html

    On laisse tomber ce qui n'est pas une page du site : le http:// non securise,
    les variantes ?lang=, un autre domaine. Ces lignes existent dans l'export mais
    ne correspondent a aucun fichier sur le disque, les compter fausserait le
    rapprochement."""
    if not url.startswith(BASE + "/"):
        return None
    reste = url[len(BASE) + 1:]
    if "?" in reste or "#" in reste:
        return None
    if reste == "":
        return "index.html"
    if reste in ("en/", "es/", "pt/", "sw/"):
        return reste + "index.html"
    return reste if reste.endswith(".html") else None


def main():
    dossier = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAUT
    if not dossier.is_dir():
        sys.exit(f"Dossier introuvable : {dossier}")

    couverture = plus_recent(dossier, "*Coverage*.xlsx")
    perf = plus_recent(dossier, "*Performance-on-Search*.xlsx")
    if not couverture and not perf:
        sys.exit(f"Aucun export trouve dans {dossier}")

    resume = {
        "genere_le": date.today().isoformat(),
        "exports": {},
        "indexation": {},
        "motifs_hors_index": {},
        "total": {},
        "pages_vues_en_recherche": {},
        "requetes_proches_page_1": [],
    }

    # --- Indexation des pages -------------------------------------------------
    if couverture:
        resume["exports"]["couverture"] = couverture.name
        resume["exports"]["couverture_date"] = date_du_nom(couverture)
        wb = openpyxl.load_workbook(couverture, read_only=True)

        # Le graphique donne l'etat jour par jour. La derniere ligne renseignee
        # est la photo la plus recente : la Search Console a deux ou trois jours
        # de retard, les dernieres lignes sont souvent vides.
        for r in reversed(lignes(onglet(wb, "Graphique"))):
            if r[1] is not None and r[2] is not None:
                resume["indexation"] = {
                    "le": str(r[0])[:10],
                    "hors_index": int(r[1]),
                    "dans_index": int(r[2]),
                }
                break

        for r in lignes(onglet(wb, "Problemes critiques")) + lignes(
            onglet(wb, "Problemes non critiques")
        ):
            if r[3] is not None:
                resume["motifs_hors_index"][str(r[0])] = int(r[3])
        wb.close()

    # --- Performances ---------------------------------------------------------
    if perf:
        resume["exports"]["performances"] = perf.name
        resume["exports"]["performances_date"] = date_du_nom(perf)
        wb = openpyxl.load_workbook(perf, read_only=True)

        # Le total se prend sur l'onglet Appareils et non sur celui des requetes :
        # la Search Console masque les requetes rares, leur somme est donc plus
        # basse que la realite. Les appareils, eux, couvrent tout le trafic.
        clics = impressions = 0
        for r in lignes(onglet(wb, "Appareils")):
            clics += int(r[1] or 0)
            impressions += int(r[2] or 0)
        if impressions:
            resume["total"] = {
                "clics": clics,
                "impressions": impressions,
                "ctr": round(clics / impressions, 5),
            }

        for r in lignes(onglet(wb, "Pages")):
            chemin = chemin_du_site(str(r[0]))
            if chemin:
                p = resume["pages_vues_en_recherche"].setdefault(
                    chemin, {"clics": 0, "impressions": 0, "position": None}
                )
                p["clics"] += int(r[1] or 0)
                p["impressions"] += int(r[2] or 0)
                # Plusieurs URLs peuvent retomber sur la meme page (/ et /index.html).
                # On garde la meilleure position, c'est celle que voit le lecteur.
                pos = round(float(r[4]), 1) if r[4] is not None else None
                if pos is not None and (p["position"] is None or pos < p["position"]):
                    p["position"] = pos

        # Une requete entre la 11e et la 20e place est a un cran de la page 1.
        # C'est le seul endroit du rapport ou un petit effort deplace vraiment
        # quelque chose, donc on la sort nommement.
        for r in lignes(onglet(wb, "Requetes")):
            pos = float(r[4]) if r[4] is not None else 999
            imp = int(r[2] or 0)
            if 10 < pos <= 20 and imp >= 10:
                resume["requetes_proches_page_1"].append(
                    {"requete": str(r[0]), "impressions": imp, "position": round(pos, 1)}
                )
        resume["requetes_proches_page_1"].sort(key=lambda x: -x["impressions"])
        wb.close()

    SORTIE.write_text(
        json.dumps(resume, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    idx = resume["indexation"]
    print(f"Ecrit {SORTIE.relative_to(ROOT)}")
    if idx:
        print(f"  indexation au {idx['le']} : {idx['dans_index']} dans l'index, "
              f"{idx['hors_index']} hors index")
    if resume["total"]:
        t = resume["total"]
        print(f"  3 mois : {t['clics']} clics, {t['impressions']} impressions, "
              f"CTR {t['ctr'] * 100:.2f} %")
    print(f"  {len(resume['pages_vues_en_recherche'])} pages vues en recherche, "
          f"{len(resume['requetes_proches_page_1'])} requetes en position 11 a 20")


if __name__ == "__main__":
    main()
